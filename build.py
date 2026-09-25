"""JSON テンプレートと素材画像から pptx を生成する。

使い方:
    uv run build.py ポスター.json

もとの pptx（"base"）を開き、JSON に書かれた要素だけを変更して "output" に保存する。
もとの pptx は上書きしない。

- image: "path" の画像で差し替える。画像は元の図形の中身だけを入れ替えるので、
  重なり順・枠線・影などの書式は残る。"path" が未指定またはファイルがない場合は、
  点線の枠を残す。
  "fit": "contain"（枠内に収める・縦横比維持）/ "cover"（枠を埋めてはみ出しをトリミング）
         / "stretch"（枠に合わせて伸縮）
- text: "text" の段落リストが元と異なる段落だけ書き換える。書き換えた段落は
  その段落の先頭の文字書式が全体に適用される。段落内改行は "\n"。
- box: [left, top, width, height]（cm）が元と異なれば移動・リサイズする。
"""

import argparse
import copy
import json
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

from template_common import (
    cm_to_emu,
    get_box,
    get_paragraphs,
    is_empty_picture_placeholder,
    iter_shapes,
)

FRAME_COLOR = RGBColor(0x99, 0x99, 0x99)


def warn(message):
    print(f"  [警告] {message}", file=sys.stderr)


# ---------------------------------------------------------------- 共通


def apply_box(shape, box):
    if box is None or box == get_box(shape):
        return
    shape.left, shape.top, shape.width, shape.height = (cm_to_emu(v) for v in box)


def drop_unused_rels(part, rIds):
    """どこからも参照されなくなったリレーションを削除し、不要な画像を保存対象から外す。"""
    used = set(part._element.xpath("//@r:embed | //@r:link | //@r:id"))
    for rId in rIds:
        if rId and rId not in used and rId in part.rels:
            part.rels.pop(rId)


# ---------------------------------------------------------------- 画像


def blip_rIds(pic_element):
    return [v for v in pic_element.xpath(".//@r:embed | .//@r:link")]


def replace_picture(shape, image_path):
    """p:pic の画像だけを差し替える（位置・書式・重なり順は維持）。"""
    pic = shape._element
    old_rIds = blip_rIds(pic)
    _, rId = shape.part.get_or_add_image_part(str(image_path))
    blip = pic.blipFill.blip
    blip.rEmbed = rId
    # SVG 画像は extLst 内の svgBlip が優先表示されるので消す
    for ext_lst in blip.findall(qn("a:extLst")):
        blip.remove(ext_lst)
    drop_unused_rels(shape.part, old_rIds)


def apply_fit(picture, image_path, fit):
    picture.crop_left = picture.crop_right = picture.crop_top = picture.crop_bottom = 0
    if fit == "stretch":
        return

    with Image.open(image_path) as img:
        img_w, img_h = img.size
    img_ratio = img_w / img_h
    left, top, width, height = picture.left, picture.top, picture.width, picture.height
    box_ratio = width / height

    if fit == "contain":
        if box_ratio > img_ratio:
            new_width = round(height * img_ratio)
            picture.left = left + (width - new_width) // 2
            picture.width = new_width
        else:
            new_height = round(width / img_ratio)
            picture.top = top + (height - new_height) // 2
            picture.height = new_height
    elif fit == "cover":
        if box_ratio > img_ratio:
            crop = (1 - img_ratio / box_ratio) / 2
            picture.crop_top = picture.crop_bottom = crop
        else:
            crop = (1 - box_ratio / img_ratio) / 2
            picture.crop_left = picture.crop_right = crop
    else:
        warn(f'fit "{fit}" は不明です（contain / cover / stretch）。stretch として扱います')


def replace_with_frame(slide, shape, label):
    """画像がない場合、画像を消して同じ位置に点線の枠を残す。"""
    frame = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, shape.left, shape.top, shape.width, shape.height)
    frame.name = shape.name
    frame.fill.background()
    frame.line.color.rgb = FRAME_COLOR
    frame.line.dash_style = MSO_LINE.DASH
    frame.line.width = Pt(1.5)
    frame.shadow.inherit = False
    tf = frame.text_frame
    tf.text = f"画像未配置\n{label}"
    font_size = Pt(max(18, min(frame.width, frame.height) / 12 / 12700))  # 枠の大きさに合わせる
    for p in tf.paragraphs:
        p.alignment = PP_ALIGN.CENTER
        for r in p.runs:
            r.font.color.rgb = FRAME_COLOR
            r.font.size = font_size

    old_rIds = blip_rIds(shape._element)
    # add_shape はスライド最前面に追加されるので、元の画像の位置（グループ内も含む）へ移す
    shape._element.addprevious(frame._element)
    shape._element.getparent().remove(shape._element)
    drop_unused_rels(slide.part, old_rIds)


def build_image(slide, shape, element, image_root):
    name = shape.name
    apply_box(shape, element.get("box"))
    path = element.get("path")
    image_path = (image_root / path) if path else None

    if image_path is None or not image_path.is_file():
        if image_path is not None:
            warn(f"{shape.name}: 画像が見つかりません: {image_path}")
        else:
            print(f"  {shape.name}: 画像未指定のため枠を残します")
        if not is_empty_picture_placeholder(shape):
            replace_with_frame(slide, shape, path or shape.name)
        return

    if is_empty_picture_placeholder(shape):
        box = (shape.left, shape.top, shape.width, shape.height)
        picture = shape.insert_picture(str(image_path))
        # 挿入直後の画像は位置をレイアウトから継承しているので、明示的に書き込む
        picture.left, picture.top, picture.width, picture.height = box
    else:
        replace_picture(shape, image_path)
        picture = shape
    apply_fit(picture, image_path, element.get("fit", "contain"))
    print(f"  {name}: {path}")


# ---------------------------------------------------------------- テキスト


def set_paragraph_text(p, text):
    """段落の文字列を置き換える。書式は段落の先頭の文字書式を引き継ぐ。"""
    first_rPr = next((c.find(qn("a:rPr")) for c in p.content_children if c.find(qn("a:rPr")) is not None), None)
    for child in p.content_children:
        p.remove(child)
    p.append_text(text)
    if first_rPr is not None:
        for r in p.r_lst:
            r.insert(0, copy.deepcopy(first_rPr))


def build_text(shape, element):
    apply_box(shape, element.get("box"))
    new_texts = element.get("text")
    if new_texts is None:
        return
    if isinstance(new_texts, str):
        new_texts = [new_texts]
    if new_texts == get_paragraphs(shape):
        return

    txBody = shape.text_frame._txBody
    ps = txBody.p_lst
    # 段落数を合わせる（増やすときは最後の段落の書式を複製）
    while len(ps) < len(new_texts):
        ps[-1].addnext(copy.deepcopy(ps[-1]))
        ps = txBody.p_lst
    for extra in ps[len(new_texts):]:
        txBody.remove(extra)
    ps = txBody.p_lst

    for p, text in zip(ps, new_texts):
        if p.text.replace("\v", "\n") != text:
            set_paragraph_text(p, text)
    print(f"  {shape.name}: テキストを更新")


# ---------------------------------------------------------------- メイン


def build(json_path):
    template = json.loads(json_path.read_text(encoding="utf-8"))
    json_dir = json_path.parent
    base_path = json_dir / template["base"]
    output_path = json_dir / template["output"]
    image_root = json_dir / template.get("image_root", ".")

    if template.get("unit", "cm") != "cm":
        sys.exit('エラー: "unit" は "cm" のみ対応しています')
    if base_path.resolve() == output_path.resolve():
        sys.exit('エラー: "output" が "base" と同じです。もとの pptx を上書きしないよう別名にしてください')

    prs = Presentation(base_path)
    for slide_spec in template["slides"]:
        index = slide_spec["index"]
        slide = prs.slides[index]
        print(f"スライド {index + 1}")
        shapes = {shape.shape_id: shape for shape, _ in iter_shapes(slide.shapes)}

        for element in slide_spec["elements"]:
            shape = shapes.get(element["id"])
            if shape is None:
                warn(f'id {element["id"]}（{element.get("name")}）の図形が見つかりません')
                continue
            if element["type"] == "image":
                build_image(slide, shape, element, image_root)
            elif element["type"] == "text":
                build_text(shape, element)
            else:
                warn(f'type "{element["type"]}" は不明です')

    output_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        prs.save(output_path)
    except PermissionError:
        sys.exit(f"エラー: {output_path} が開いたままです。閉じてから再実行してください")
    print(f"\n{output_path} に保存しました")


def main():
    parser = argparse.ArgumentParser(description="JSON テンプレートから pptx を生成する")
    parser.add_argument("json", type=Path, help="extract.py で作った JSON テンプレート")
    args = parser.parse_args()
    build(args.json.resolve())


if __name__ == "__main__":
    main()
