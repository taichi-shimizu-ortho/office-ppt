"""既存の pptx から、画像枠とテキストを JSON テンプレートとして書き出す。

使い方:
    uv run extract.py ポスター.pptx [-o ポスター.json]

画像はすべて差し替える前提なので中身は書き出さず、枠（位置・サイズ）だけを
"path": null のプレースホルダーとして出力する。テキストは段落ごとの文字列。
"""

import argparse
import json
import os
import re
from pathlib import Path

from pptx import Presentation

from template_common import emu_to_cm, get_box, get_paragraphs, is_image, iter_shapes


def extract(pptx_path, json_path):
    prs = Presentation(pptx_path)
    json_dir = json_path.parent

    slides = []
    for index, slide in enumerate(prs.slides):
        elements = []
        for shape, group in iter_shapes(slide.shapes):
            if is_image(shape):
                element = {
                    "id": shape.shape_id,
                    "name": shape.name,
                    "type": "image",
                    "box": get_box(shape),
                    "path": None,
                    "fit": "contain",
                }
            elif shape.has_text_frame and shape.text_frame.text.strip():
                element = {
                    "id": shape.shape_id,
                    "name": shape.name,
                    "type": "text",
                    "box": get_box(shape),
                    "text": get_paragraphs(shape),
                }
            else:
                continue
            if group:
                element["group"] = group
            elements.append(element)
        slides.append({"index": index, "elements": elements})

    template = {
        "base": os.path.relpath(pptx_path, json_dir),
        "output": os.path.relpath(pptx_path.with_name(pptx_path.stem + "_out.pptx"), json_dir),
        "image_root": ".",
        "unit": "cm",
        "slide_size": [emu_to_cm(prs.slide_width), emu_to_cm(prs.slide_height)],
        "slides": slides,
    }
    text = json.dumps(template, ensure_ascii=False, indent=2)
    # 数値の配列（box など）は1行にまとめて読みやすくする
    text = re.sub(r"\[\s*([-\d.,\s]+?)\s*\]", lambda m: "[" + re.sub(r",\s*", ", ", m.group(1)) + "]", text)
    json_path.write_text(text + "\n", encoding="utf-8")

    n_images = sum(e["type"] == "image" for s in slides for e in s["elements"])
    n_texts = sum(e["type"] == "text" for s in slides for e in s["elements"])
    print(f"{json_path} を書き出しました（スライド {len(slides)} 枚、画像枠 {n_images} 個、テキスト {n_texts} 個）")


def main():
    parser = argparse.ArgumentParser(description="pptx から JSON テンプレートを書き出す")
    parser.add_argument("pptx", type=Path, help="もとにする pptx")
    parser.add_argument("-o", "--output", type=Path, help="出力する JSON（省略時は pptx と同名の .json）")
    args = parser.parse_args()

    pptx_path = args.pptx.resolve()
    json_path = (args.output or args.pptx.with_suffix(".json")).resolve()
    extract(pptx_path, json_path)


if __name__ == "__main__":
    main()
