"""extract.py / build.py で共有するユーティリティ。"""

from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.shapes.picture import Picture
from pptx.util import Emu

EMU_PER_CM = 360000


def emu_to_cm(value):
    return round(Emu(value) / EMU_PER_CM, 3)


def cm_to_emu(value):
    return Emu(round(value * EMU_PER_CM))


def iter_shapes(shapes, group=None):
    """スライド上の図形をグループの中まで再帰的に列挙する。(shape, 親グループ名) を返す。"""
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(shape.shapes, shape.name)
        else:
            yield shape, group


def is_empty_picture_placeholder(shape):
    return (
        shape.is_placeholder
        and not isinstance(shape, Picture)
        and shape.placeholder_format.type == PP_PLACEHOLDER.PICTURE
    )


def is_image(shape):
    return isinstance(shape, Picture) or is_empty_picture_placeholder(shape)


def get_box(shape):
    """[left, top, width, height] を cm で返す。グループ内の図形はグループ内座標。"""
    return [emu_to_cm(v) for v in (shape.left, shape.top, shape.width, shape.height)]


def get_paragraphs(shape):
    # 段落内改行（Shift+Enter）は "\v" で返るので、JSON上は "\n" で表す
    return [p.text.replace("\v", "\n") for p in shape.text_frame.paragraphs]
